# -*- coding: utf-8 -*-
"""把 docs/与天津大学脑机方向技术交流问题清单.md 转成 Word 文档。

支持：标题(#/##/###)、表格、> 引用、有序/无序列表(含缩进)、---、复选框、
行内 **粗体**、`代码`、*斜体*。
"""
import re
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, Cm, RGBColor

SRC = r"D:\校企合作项目\docs\与天津大学脑机方向技术交流问题清单.md"
OUT = r"D:\校企合作项目\与天大脑机方向技术交流问题清单.docx"

BODY_FONT = "宋体"
HEAD_FONT = "黑体"
MONO_FONT = "Consolas"
CONTENT_WIDTH_CM = 16.6

doc = Document()
normal = doc.styles["Normal"]
normal.font.name = BODY_FONT
normal.font.size = Pt(10.5)
normal._element.rPr.rFonts.set(qn("w:eastAsia"), BODY_FONT)
normal.paragraph_format.space_after = Pt(4)
normal.paragraph_format.line_spacing = 1.25

for name, size in (("Heading 1", 15), ("Heading 2", 12.5), ("Heading 3", 11.5)):
    st = doc.styles[name]
    st.font.name = HEAD_FONT
    st.font.size = Pt(size)
    st.font.color.rgb = RGBColor(0x1F, 0x38, 0x64)
    st._element.rPr.rFonts.set(qn("w:eastAsia"), HEAD_FONT)

for section in doc.sections:
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.left_margin = Cm(2.2)
    section.right_margin = Cm(2.2)
    section.top_margin = Cm(2.0)
    section.bottom_margin = Cm(2.0)


def style_run(run, bold=False, italic=False, code=False):
    if code:
        run.font.name = MONO_FONT
        run._element.rPr.rFonts.set(qn("w:eastAsia"), MONO_FONT)
        run.font.size = Pt(9.5)
        run.font.color.rgb = RGBColor(0xA3, 0x1D, 0x1D)
    else:
        run.font.name = BODY_FONT
        run._element.rPr.rFonts.set(qn("w:eastAsia"), BODY_FONT)
    run.font.bold = bold
    run.font.italic = italic
    return run


TOKEN = re.compile(r"(\*\*.+?\*\*|`[^`]+`|(?<!\*)\*(?!\*)[^*]+\*(?!\*))")


def add_inline(paragraph, text):
    text = text.replace("\\|", "|")
    for piece in TOKEN.split(text):
        if not piece:
            continue
        if piece.startswith("**") and piece.endswith("**") and len(piece) > 4:
            style_run(paragraph.add_run(piece[2:-2]), bold=True)
        elif piece.startswith("`") and piece.endswith("`") and len(piece) > 2:
            style_run(paragraph.add_run(piece[1:-1]), code=True)
        elif piece.startswith("*") and piece.endswith("*") and len(piece) > 2:
            style_run(paragraph.add_run(piece[1:-1]), italic=True)
        else:
            style_run(paragraph.add_run(piece))


def shade(cell, hexcolor):
    tcPr = cell._tc.get_or_add_tcPr()
    el = OxmlElement("w:shd")
    el.set(qn("w:val"), "clear")
    el.set(qn("w:color"), "auto")
    el.set(qn("w:fill"), hexcolor)
    tcPr.append(el)


def split_row(line):
    line = line.strip()
    if line.startswith("|"):
        line = line[1:]
    if line.endswith("|"):
        line = line[:-1]
    return [c.strip() for c in line.split("|")]


def force_fixed_layout(t, total_cm, widths):
    """锁定固定布局、总宽，并让 tblGrid 与单元格宽度一致（否则渲染器按网格撑宽表格）。"""
    tblPr = t._tbl.tblPr
    for old in tblPr.findall(qn("w:tblLayout")):
        tblPr.remove(old)
    layout = OxmlElement("w:tblLayout")
    layout.set(qn("w:type"), "fixed")
    tblPr.append(layout)
    tblW = tblPr.find(qn("w:tblW"))
    if tblW is None:
        tblW = OxmlElement("w:tblW")
        tblPr.append(tblW)
    tblW.set(qn("w:type"), "dxa")
    tblW.set(qn("w:w"), str(int(round(total_cm * 567))))
    grid = t._tbl.find(qn("w:tblGrid"))
    if grid is not None:
        for col, w in zip(grid.findall(qn("w:gridCol")), widths):
            col.set(qn("w:w"), str(int(round(w.cm * 567))))


def add_doc_table(rows):
    ncol = len(rows[0])
    t = doc.add_table(rows=0, cols=ncol)
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    for i, row in enumerate(rows):
        cells = t.add_row().cells
        for j in range(ncol):
            val = row[j] if j < len(row) else ""
            cell = cells[j]
            cell.text = ""
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(1)
            p.paragraph_format.line_spacing = 1.1
            add_inline(p, val)
            for r in p.runs:
                r.font.size = Pt(9)
                if i == 0:
                    r.font.bold = True
            if i == 0:
                shade(cell, "DCE6F1")
    # 列宽按内容长度分配，限幅后归一化
    weights = []
    for j in range(ncol):
        lens = [len(r[j]) if j < len(r) else 0 for r in rows]
        avg = sum(lens) / max(len(lens), 1)
        weights.append(max(4.0, min(avg, 26.0)))
    total = sum(weights)
    widths = [Cm(CONTENT_WIDTH_CM * w / total) for w in weights]
    for r in t.rows:
        for j, w in enumerate(widths):
            r.cells[j].width = w
    force_fixed_layout(t, CONTENT_WIDTH_CM, widths)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return t


def add_code_line(text, first=False, last=False):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.left_indent = Cm(0.6)
    pf.right_indent = Cm(0.2)
    pf.space_before = Pt(4 if first else 0)
    pf.space_after = Pt(4 if last else 0)
    pf.line_spacing = 1.0
    run = p.add_run(text)
    run.font.name = MONO_FONT
    run._element.rPr.rFonts.set(qn("w:eastAsia"), MONO_FONT)
    run.font.size = Pt(8.5)
    run.font.color.rgb = RGBColor(0x33, 0x33, 0x33)
    pPr = p._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), "F2F2F2")
    pPr.append(shd)
    return p


def flush_quote(buf):
    if not buf:
        return
    for item in buf:
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Cm(0.5)
        p.paragraph_format.space_after = Pt(2)
        add_inline(p, item)
        for r in p.runs:
            r.font.color.rgb = RGBColor(0x40, 0x40, 0x40)
            r.font.size = Pt(10)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    buf.clear()


lines = open(SRC, encoding="utf-8").read().splitlines()
i = 0
quote_buf = []
while i < len(lines):
    raw = lines[i]
    line = raw.rstrip()
    stripped = line.strip()

    if not stripped:
        i += 1
        continue

    # 水平线
    if re.fullmatch(r"-{3,}|\*{3,}", stripped):
        flush_quote(quote_buf)
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(2)
        style_run(p.add_run("─" * 46), bold=False)
        for r in p.runs:
            r.font.color.rgb = RGBColor(0xAA, 0xAA, 0xAA)
            r.font.size = Pt(8)
        i += 1
        continue

    # 围栏代码块
    if stripped.startswith("```"):
        flush_quote(quote_buf)
        i += 1
        buf = []
        while i < len(lines) and not lines[i].strip().startswith("```"):
            buf.append(lines[i].rstrip())
            i += 1
        i += 1
        for k, cl in enumerate(buf):
            add_code_line(cl, first=(k == 0), last=(k == len(buf) - 1))
        doc.add_paragraph().paragraph_format.space_after = Pt(2)
        continue

    # 表格
    if stripped.startswith("|"):
        flush_quote(quote_buf)
        block = []
        while i < len(lines) and lines[i].strip().startswith("|"):
            block.append(lines[i])
            i += 1
        rows = []
        for k, b in enumerate(block):
            cells = split_row(b)
            if k == 1 and all(re.fullmatch(r":?-{2,}:?", c or "-") for c in cells):
                continue
            rows.append(cells)
        if rows:
            add_doc_table(rows)
        continue

    # 引用
    if stripped.startswith(">"):
        quote_buf.append(stripped.lstrip(">").strip())
        i += 1
        continue
    flush_quote(quote_buf)

    m = re.match(r"^(#{1,6})\s+(.*)$", stripped)
    if m:
        level = len(m.group(1))
        text = m.group(2).strip()
        p = doc.add_paragraph(style=f"Heading {min(level, 3)}")
        p.paragraph_format.space_before = Pt(10 if level <= 2 else 6)
        p.paragraph_format.space_after = Pt(4)
        add_inline(p, text)
        for r in p.runs:
            r.font.name = HEAD_FONT
            r._element.rPr.rFonts.set(qn("w:eastAsia"), HEAD_FONT)
            r.font.color.rgb = RGBColor(0x1F, 0x38, 0x64)
            r.font.bold = True
        i += 1
        continue

    # 列表
    m = re.match(r"^(\s*)([-*+]|\d+\.)\s+(.*)$", line)
    if m:
        indent = len(m.group(1))
        marker = m.group(2)
        text = m.group(3).strip()
        text = re.sub(r"^\[( |x|X)\]\s*", "□ ", text)
        bullet = "•" if marker in "-*+" else marker
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Cm(0.45 + 0.45 * (indent // 2))
        p.paragraph_format.first_line_indent = Cm(-0.45) if marker in "-*+" else Cm(-0.6)
        p.paragraph_format.space_after = Pt(2)
        add_inline(p, f"{bullet} {text}" if marker in "-*+" else f"{bullet} {text}")
        i += 1
        continue

    # 普通段落
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(4)
    add_inline(p, stripped)
    i += 1

flush_quote(quote_buf)
doc.save(OUT)
print("saved:", OUT)
