"""把 docs/ 下的 Markdown 文档转换为 Word（.docx），用于对外提交。

用法：
    python .tools/md2docx.py <输入.md> [输出.docx] [--footer "页脚文字"]
    python .tools/md2docx.py --all          # 批量转换 docs/ 下的规格文档

保留结构：标题层级、正文、表格、代码块、列表（含 [ ] 任务框）、引用块。
中文需显式设置 w:eastAsia 字体，否则西文字体名无法覆盖中文。
表格必须同时写 w:tblW / w:gridCol / w:tcW，否则 LibreOffice 与 Word 会按内容撑宽导致溢出。

（本脚本由 `.tools/_md2docx_from_bci-decoder.py`（原 `bci-decoder/_md2docx.py`）泛化而来：改为按参数接收输入输出，便于批量转文档。）
"""

from __future__ import annotations

import argparse
import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

BODY_EA = "微软雅黑"
BODY_LATIN = "Segoe UI"
MONO = "Consolas"
ACCENT = RGBColor(0x1F, 0x45, 0x7A)
TWIPS_PER_CM = 567
TEXT_WIDTH_CM = 16.3  # A4 版心 16.60cm，留余量避免贴边


# ---------- 底层工具 ----------

def set_run_fonts(run, latin=BODY_LATIN, ea=BODY_EA, size=None, bold=None,
                  color=None, mono=False):
    font = run.font
    if mono:
        latin = ea = MONO
    font.name = latin
    rpr = run._element.get_or_add_rPr()
    rf = rpr.get_or_add_rFonts()
    rf.set(qn("w:ascii"), latin)
    rf.set(qn("w:hAnsi"), latin)
    rf.set(qn("w:eastAsia"), ea)
    rf.set(qn("w:cs"), latin)
    if size is not None:
        font.size = Pt(size)
    if bold is not None:
        font.bold = bold
    if color is not None:
        font.color.rgb = color


def shade(element, fill: str):
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    element.append(shd)


def para_shading(paragraph, fill: str):
    shade(paragraph._p.get_or_add_pPr(), fill)


def cell_shading(cell, fill: str):
    shade(cell._tc.get_or_add_tcPr(), fill)


def set_repeat_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    el = OxmlElement("w:tblHeader")
    el.set(qn("w:val"), "true")
    tr_pr.append(el)


INLINE_RE = re.compile(r"(\*\*.+?\*\*|`[^`]+`|\[[^\]]+\]\([^)]+\))")
LINK_RE = re.compile(r"\[([^\]]+)\]\(([^)]+)\)")


def add_inline(paragraph, text: str, size=10.5, base_bold=False, color=None):
    """解析 **粗体** / `代码` / [链接](url)。"""
    for token in INLINE_RE.split(text):
        if not token:
            continue
        if token.startswith("**") and token.endswith("**") and len(token) > 4:
            r = paragraph.add_run(token[2:-2])
            set_run_fonts(r, size=size, bold=True, color=color)
        elif token.startswith("`") and token.endswith("`") and len(token) > 2:
            r = paragraph.add_run(token[1:-1])
            set_run_fonts(r, size=size - 0.5, mono=True, color=RGBColor(0xA3, 0x1D, 0x1D))
        else:
            m = LINK_RE.match(token)
            if m:
                r = paragraph.add_run(m.group(1))
                set_run_fonts(r, size=size, bold=base_bold,
                              color=RGBColor(0x0B, 0x5C, 0xAB))
                r.font.underline = True
            else:
                r = paragraph.add_run(token)
                set_run_fonts(r, size=size, bold=base_bold, color=color)


# ---------- 文档基础设置 ----------

def setup_styles(doc: Document):
    normal = doc.styles["Normal"]
    normal.font.name = BODY_LATIN
    normal.font.size = Pt(10.5)
    rpr = normal.element.get_or_add_rPr()
    rf = rpr.get_or_add_rFonts()
    rf.set(qn("w:ascii"), BODY_LATIN)
    rf.set(qn("w:hAnsi"), BODY_LATIN)
    rf.set(qn("w:eastAsia"), BODY_EA)
    normal.paragraph_format.space_after = Pt(4)
    normal.paragraph_format.line_spacing = 1.25

    for name, size in (("Heading 1", 16), ("Heading 2", 13), ("Heading 3", 11.5)):
        st = doc.styles[name]
        st.font.name = BODY_LATIN
        st.font.size = Pt(size)
        st.font.bold = True
        st.font.color.rgb = ACCENT
        st.element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), BODY_EA)
        st.paragraph_format.space_before = Pt(12)
        st.paragraph_format.space_after = Pt(6)


def add_footer(doc: Document, text: str):
    p = doc.sections[0].footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(text)
    set_run_fonts(r, size=8, color=RGBColor(0x80, 0x80, 0x80))


# ---------- 块级解析 ----------

def is_table_sep(line: str) -> bool:
    return bool(re.fullmatch(r"\|[\s:|-]+\|", line.strip())) and "-" in line


def split_row(line: str) -> list[str]:
    return [c.strip() for c in line.strip().strip("|").split("|")]


def calc_widths(rows: list[list[str]], n_cols: int, indent_cm: float = 0.3) -> list[float]:
    """按各列内容量比例分配列宽（中文按 2 个字符宽计）。"""
    def w(s: str) -> int:
        return sum(2 if ord(ch) > 0x2E80 else 1 for ch in s)

    avail = TEXT_WIDTH_CM - indent_cm
    col_w = [max(w(r[j]) for r in rows) for j in range(n_cols)]
    total = sum(col_w) or 1
    lo, hi = 2.0, 9.5
    widths = [min(max(avail * c / total, lo), hi) for c in col_w]
    for _ in range(6):
        s = sum(widths)
        if abs(s - avail) < 0.05:
            break
        scale = avail / s
        widths = [min(max(x * scale, lo), hi) for x in widths]
    return widths


def add_table(doc: Document, rows: list[list[str]]):
    n_cols = max(len(r) for r in rows)
    rows = [r + [""] * (n_cols - len(r)) for r in rows]
    table = doc.add_table(rows=0, cols=n_cols)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    widths = calc_widths(rows, n_cols)
    for col, width in zip(table.columns, widths):
        col.width = Cm(width)

    for i, row_cells in enumerate(rows):
        cells = table.add_row().cells
        for j, txt in enumerate(row_cells):
            cell = cells[j]
            cell.width = Cm(widths[j])
            # 必须显式写 w:tcW，否则单元格按内容自动变宽，撑破页面
            tc_pr = cell._tc.get_or_add_tcPr()
            for el in tc_pr.findall(qn("w:tcW")):
                tc_pr.remove(el)
            tcw = OxmlElement("w:tcW")
            tcw.set(qn("w:type"), "dxa")
            tcw.set(qn("w:w"), str(int(round(widths[j] * TWIPS_PER_CM))))
            tc_pr.append(tcw)
            cell.text = ""
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.line_spacing = 1.1
            if i == 0:
                add_inline(p, txt, size=10, base_bold=True)
                cell_shading(cell, "DCE6F1")
            else:
                add_inline(p, txt, size=10)

    tbl_pr = table._tbl.tblPr
    for tag in ("w:tblLayout", "w:tblW"):
        for old in tbl_pr.findall(qn(tag)):
            tbl_pr.remove(old)
    layout = OxmlElement("w:tblLayout")
    layout.set(qn("w:type"), "fixed")
    tbl_pr.append(layout)
    # 显式表宽：w:tblW=auto/0 会让列宽被忽略
    tbl_w = OxmlElement("w:tblW")
    tbl_w.set(qn("w:type"), "dxa")
    tbl_w.set(qn("w:w"), str(int(round(sum(widths) * TWIPS_PER_CM))))
    tbl_pr.append(tbl_w)

    set_repeat_header(table.rows[0])
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return table


def add_code_block(doc: Document, lines: list[str]):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.space_before = Pt(4)
    pf.space_after = Pt(8)
    pf.left_indent = Cm(0.4)
    pf.line_spacing = 1.0
    para_shading(p, "F5F5F5")
    for k, line in enumerate(lines):
        r = p.add_run(line)
        set_run_fonts(r, size=9, mono=True)
        if k < len(lines) - 1:
            r.add_break()
    return p


def add_quote(doc: Document, lines: list[str]):
    joined = " ".join(l.strip() for l in lines).strip()
    strong = joined.startswith("**")
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.left_indent = Cm(0.5)
    pf.space_before = Pt(6)
    pf.space_after = Pt(8)
    if strong:
        para_shading(p, "FFF4D6")
        pf.left_indent = Cm(0.3)
    for token in INLINE_RE.split(joined):
        if not token:
            continue
        if token.startswith("**") and token.endswith("**") and len(token) > 4:
            r = p.add_run(token[2:-2])
            set_run_fonts(r, size=11 if strong else 10.5, bold=True,
                          color=ACCENT if strong else None)
        elif token.startswith("`") and token.endswith("`") and len(token) > 2:
            r = p.add_run(token[1:-1])
            set_run_fonts(r, size=9.5, mono=True, color=RGBColor(0xA3, 0x1D, 0x1D))
        else:
            r = p.add_run(token)
            set_run_fonts(r, size=11 if strong else 10.5, bold=strong,
                          color=ACCENT if strong else None)


def add_list_item(doc: Document, line: str, num_state: dict):
    m = re.match(r"^(\s*)([-*]|\d+\.)\s+(.*)$", line)
    indent, marker, text = m.group(1), m.group(2), m.group(3)
    depth = len(indent) // 2
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.space_after = Pt(2)
    pf.line_spacing = 1.2
    pf.left_indent = Cm(0.75 + 0.55 * depth)
    pf.first_line_indent = Cm(-0.55)

    if marker in ("-", "*"):
        num_state["n"] = 0
        # 任务清单：[ ] / [x] → 勾选框
        mm = re.match(r"^\[([ xX])\]\s*(.*)$", text)
        if mm:
            checked = mm.group(1).lower() == "x"
            text = mm.group(2)
            r = p.add_run(("\u2611 " if checked else "\u2610 ") + "\t")
            set_run_fonts(r, size=10.5)
        else:
            bullet = "\u2022" if depth == 0 else "\u25e6"
            r = p.add_run(bullet + "\t")
            set_run_fonts(r, size=10.5)
    else:
        num_state["n"] = num_state.get("n", 0) + 1
        r = p.add_run(f"{num_state['n']}.\t")
        set_run_fonts(r, size=10.5)
    add_inline(p, text, size=10.5)


# ---------- 主流程 ----------

def convert(src: Path, out: Path, footer: str) -> int:
    text = src.read_text(encoding="utf-8")
    lines = text.splitlines()

    doc = Document()
    setup_styles(doc)
    sec = doc.sections[0]
    sec.page_width = Cm(21.0)      # A4
    sec.page_height = Cm(29.7)
    sec.left_margin = sec.right_margin = Cm(2.2)
    sec.top_margin = sec.bottom_margin = Cm(2.0)
    add_footer(doc, footer)

    i = 0
    first_h1_done = False
    in_code = False
    code_buf: list[str] = []
    quote_buf: list[str] = []
    num_state: dict = {"n": 0}

    def flush_quote():
        nonlocal quote_buf
        if quote_buf:
            add_quote(doc, quote_buf)
            quote_buf = []

    while i < len(lines):
        raw = lines[i]
        line = raw.rstrip()

        if line.strip().startswith("```"):
            if in_code:
                add_code_block(doc, code_buf)
                code_buf = []
                in_code = False
            else:
                flush_quote()
                in_code = True
            i += 1
            continue
        if in_code:
            code_buf.append(raw)
            i += 1
            continue

        if line.lstrip().startswith(">"):
            quote_buf.append(line.lstrip()[1:].strip())
            i += 1
            continue
        flush_quote()

        if re.fullmatch(r"-{3,}", line.strip()):
            p = doc.add_paragraph()
            p.paragraph_format.space_after = Pt(2)
            pPr = p._p.get_or_add_pPr()
            bd = OxmlElement("w:pBdr")
            bottom = OxmlElement("w:bottom")
            bottom.set(qn("w:val"), "single")
            bottom.set(qn("w:sz"), "6")
            bottom.set(qn("w:color"), "BFBFBF")
            bd.append(bottom)
            pPr.append(bd)
            i += 1
            continue

        if not line.strip():
            i += 1
            continue

        m = re.match(r"^(#{1,6})\s+(.*)$", line)
        if m:
            level, title = len(m.group(1)), m.group(2).strip()
            if level == 1 and not first_h1_done:
                p = doc.add_paragraph()
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                r = p.add_run(title)
                set_run_fonts(r, size=22, bold=True, color=ACCENT)
                p.paragraph_format.space_after = Pt(4)
                first_h1_done = True
            else:
                doc.add_heading(title, level=min(level, 4))
            i += 1
            continue

        if line.lstrip().startswith("|") and i + 1 < len(lines) and is_table_sep(lines[i + 1]):
            rows = [split_row(line)]
            i += 2
            while i < len(lines) and lines[i].lstrip().startswith("|"):
                rows.append(split_row(lines[i]))
                i += 1
            add_table(doc, rows)
            continue

        if re.match(r"^\s*([-*]|\d+\.)\s+", raw):
            add_list_item(doc, raw, num_state)
            i += 1
            continue

        num_state["n"] = 0
        p = doc.add_paragraph()
        add_inline(p, line.strip())
        i += 1

    if in_code and code_buf:
        add_code_block(doc, code_buf)
    flush_quote()

    out.parent.mkdir(parents=True, exist_ok=True)
    doc.save(out)
    print(f"[ok] {src.name} -> {out.name}  段落={len(doc.paragraphs)} 表格={len(doc.tables)}")
    return 0


BATCH = [
    ("MVP规格说明书.md", "MVP规格说明书 v1.0"),
    ("需求确认记录.md", "需求确认记录 v1.0"),
    ("项目章程.md", "项目章程 v1.0"),
]


def main(argv: list[str]) -> int:
    ap = argparse.ArgumentParser(description="Markdown -> Word（中文排版）")
    ap.add_argument("src", nargs="?", help="输入 .md 路径")
    ap.add_argument("out", nargs="?", help="输出 .docx 路径（默认同名同目录）")
    ap.add_argument("--footer", default="", help="页脚文字")
    ap.add_argument("--all", action="store_true", help="批量转换 docs/ 下的规格文档")
    args = ap.parse_args(argv)

    if args.all:
        docs = Path(__file__).resolve().parent.parent / "docs"
        for name, footer in BATCH:
            src = docs / name
            if not src.exists():
                print(f"[skip] 未找到 {src}")
                continue
            convert(src, src.with_suffix(".docx"), footer)
        return 0

    if not args.src:
        ap.error("需要输入 .md 路径，或使用 --all")
    src = Path(args.src).resolve()
    out = Path(args.out).resolve() if args.out else src.with_suffix(".docx")
    footer = args.footer or f"{src.stem} · 脑控游戏运动想象在线解码模块"
    return convert(src, out, footer)


if __name__ == "__main__":
    sys.exit(main(sys.argv[1:]))
